import os
import time
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_light_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    blank_layout = prs.slide_layouts[6]
    
    # Google Cloud Inspired Light Theme Colors
    BG_LIGHT = RGBColor(248, 249, 250)      # Light gray/white background
    TEXT_DARK = RGBColor(32, 33, 36)        # Google dark gray
    TEXT_MUTED = RGBColor(95, 99, 104)      # Google secondary text
    
    # Google Brand Colors
    G_BLUE = RGBColor(66, 133, 244)
    G_RED = RGBColor(234, 67, 53)
    G_YELLOW = RGBColor(251, 188, 5)
    G_GREEN = RGBColor(52, 168, 83)
    WHITE = RGBColor(255, 255, 255)
    
    def set_bg(slide):
        bg = slide.background
        fill = bg.fill
        fill.solid()
        fill.fore_color.rgb = BG_LIGHT

        logo_wm_path = os.path.join(os.path.dirname(__file__), 'static', 'brics_logo_watermark.png')
        if os.path.exists(logo_wm_path):
            wm_width = Inches(5.6)
            wm_height = Inches(6.15)
            wm_left = (prs.slide_width - wm_width) / 2
            wm_top = (prs.slide_height - wm_height) / 2
            slide.shapes.add_picture(logo_wm_path, wm_left, wm_top, wm_width, wm_height)

    def add_header(slide, title_text, subtitle_text, slide_num):
        # Top color accent bar
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.15))
        shape.fill.solid()
        shape.fill.fore_color.rgb = G_BLUE
        shape.line.fill.background()
        
        tb = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(10.5), Inches(1.1))
        tf = tb.text_frame
        tf.word_wrap = True
        
        p = tf.paragraphs[0]
        p.text = title_text
        p.font.name = 'Arial'
        p.font.size = Pt(28)
        p.font.bold = True
        p.font.color.rgb = TEXT_DARK
        
        p2 = tf.add_paragraph()
        p2.text = subtitle_text
        p2.font.name = 'Arial'
        p2.font.size = Pt(14)
        p2.font.color.rgb = G_BLUE
        p2.space_before = Pt(4)
        
        # Slide number
        tb_num = slide.shapes.add_textbox(Inches(11.5), Inches(0.4), Inches(1.0), Inches(0.5))
        tf_num = tb_num.text_frame
        p_num = tf_num.paragraphs[0]
        p_num.text = slide_num
        p_num.font.name = 'Arial'
        p_num.font.size = Pt(12)
        p_num.font.color.rgb = TEXT_MUTED
        p_num.alignment = PP_ALIGN.RIGHT

    def add_feature_card(slide, left, top, width, height, title, body, accent=G_BLUE):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = WHITE
        shape.line.color.rgb = RGBColor(218, 220, 224) # light border
        shape.line.width = Pt(1)
        
        # Accent bar on the left of the card
        acc = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, Inches(0.15), height)
        acc.fill.solid()
        acc.fill.fore_color.rgb = accent
        acc.line.fill.background()
        
        tb = slide.shapes.add_textbox(left + Inches(0.3), top + Inches(0.2), width - Inches(0.5), height - Inches(0.4))
        tf = tb.text_frame
        tf.word_wrap = True
        
        p = tf.paragraphs[0]
        p.text = title
        p.font.name = 'Arial'
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = TEXT_DARK
        
        p2 = tf.add_paragraph()
        p2.text = body
        p2.font.name = 'Arial'
        p2.font.size = Pt(12)
        p2.font.color.rgb = TEXT_MUTED
        p2.space_before = Pt(8)

    def add_screenshot_placeholder(slide, left, top, width, height, title_str, url_str):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(240, 244, 248)
        shape.line.color.rgb = G_BLUE
        shape.line.width = Pt(2)
        
        tb = slide.shapes.add_textbox(left + Inches(0.3), top + Inches(0.3), width - Inches(0.6), height - Inches(0.6))
        tf = tb.text_frame
        tf.word_wrap = True
        
        p_head = tf.paragraphs[0]
        p_head.text = f"📸  {title_str}"
        p_head.font.name = 'Arial'
        p_head.font.size = Pt(15)
        p_head.font.bold = True
        p_head.font.color.rgb = G_BLUE
        p_head.alignment = PP_ALIGN.CENTER
        
        p_url = tf.add_paragraph()
        p_url.text = f"URL: {url_str}"
        p_url.font.name = 'Consolas'
        p_url.font.size = Pt(11)
        p_url.font.color.rgb = TEXT_MUTED
        p_url.space_before = Pt(6)
        p_url.alignment = PP_ALIGN.CENTER

    # ---------------------------------------------------------
    # SLIDE 1: TITLE
    # ---------------------------------------------------------
    s1 = prs.slides.add_slide(blank_layout)
    set_bg(s1)
    
    tb = s1.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(11.333), Inches(4.5))
    tf = tb.text_frame
    
    p = tf.paragraphs[0]
    p.text = "Build with AI: Code for Communities"
    p.font.name = 'Arial'
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = G_BLUE
    p.alignment = PP_ALIGN.CENTER
    
    p2 = tf.add_paragraph()
    p2.text = "CivicPulse-BRICS"
    p2.font.name = 'Arial'
    p2.font.size = Pt(52)
    p2.font.bold = True
    p2.font.color.rgb = TEXT_DARK
    p2.space_before = Pt(10)
    p2.alignment = PP_ALIGN.CENTER
    
    p3 = tf.add_paragraph()
    p3.text = "AI-Powered Infrastructure Governance & Citizen Triage"
    p3.font.name = 'Arial'
    p3.font.size = Pt(20)
    p3.font.color.rgb = TEXT_MUTED
    p3.space_before = Pt(14)
    p3.alignment = PP_ALIGN.CENTER
    
    p4 = tf.add_paragraph()
    p4.text = "By Team Syntax Error"
    p4.font.name = 'Arial'
    p4.font.size = Pt(16)
    p4.font.bold = True
    p4.font.color.rgb = G_GREEN
    p4.space_before = Pt(30)
    p4.alignment = PP_ALIGN.CENTER

    # ---------------------------------------------------------
    # SLIDE 2: THE PROBLEM
    # ---------------------------------------------------------
    s2 = prs.slides.add_slide(blank_layout)
    set_bg(s2)
    add_header(s2, "The Problem", "Fragmented feedback & blind spots in municipal infrastructure", "02 / 12")
    
    add_feature_card(s2, Inches(1.5), Inches(2.0), Inches(4.8), Inches(2.0), "Fragmented Citizen Reporting", "Citizens have no unified, multilingual platform to report broken infrastructure (potholes, water lines). Reports get lost in bureaucracy.", G_RED)
    add_feature_card(s2, Inches(7.0), Inches(2.0), Inches(4.8), Inches(2.0), "Manual Triage Bottlenecks", "City planners receive unstructured text and unverified photos, lacking the resources to accurately prioritize critical hazards.", G_YELLOW)
    add_feature_card(s2, Inches(1.5), Inches(4.5), Inches(4.8), Inches(2.0), "Misaligned Capital Spending", "Budgets are allocated without correlating real-time citizen demand, leading to neglected hotspots and inefficient resource use.", G_BLUE)
    add_feature_card(s2, Inches(7.0), Inches(4.5), Inches(4.8), Inches(2.0), "Language & Accessibility Barriers", "Diverse populations in BRICS nations are excluded from participatory budgeting due to language barriers and complex apps.", G_GREEN)

    # ---------------------------------------------------------
    # SLIDE 3: THE SOLUTION
    # ---------------------------------------------------------
    s3 = prs.slides.add_slide(blank_layout)
    set_bg(s3)
    add_header(s3, "The Solution: CivicPulse-BRICS", "A sovereign Digital Public Good for infrastructure planning", "03 / 12")
    
    add_feature_card(s3, Inches(1.0), Inches(2.5), Inches(3.5), Inches(3.5), "1. Multimodal Citizen App", "Citizens can report issues using photos, voice notes, and GPS. Community demands can be upvoted for participatory budgeting.", G_BLUE)
    add_feature_card(s3, Inches(4.9), Inches(2.5), Inches(3.5), Inches(3.5), "2. AI Defect Triage", "Our backend instantly analyzes uploaded photos and voice notes using AI to assess severity, defect type, and structural hazard.", G_RED)
    add_feature_card(s3, Inches(8.8), Inches(2.5), Inches(3.5), Inches(3.5), "3. Official Command Center", "Provides city officials with an auto-prioritized repair queue, GIS heatmaps, and a CapEx budget misalignment matrix.", G_GREEN)

    # ---------------------------------------------------------
    # SLIDE 4: AI APPROACH
    # ---------------------------------------------------------
    s4 = prs.slides.add_slide(blank_layout)
    set_bg(s4)
    add_header(s4, "Our AI Approach (Google Gemini)", "How we process unstructured citizen petitions into actionable data", "04 / 12")
    
    add_feature_card(s4, Inches(1.5), Inches(2.0), Inches(10.3), Inches(1.8), "Gemini Vision: Infrastructure Defect Analysis", "When a citizen uploads a photo of broken infrastructure, we prompt Google Gemini to analyze the image. It returns a structured JSON assessing the defect category, severity score (1-10), and potential public risk.", G_BLUE)
    add_feature_card(s4, Inches(1.5), Inches(4.2), Inches(10.3), Inches(1.8), "Gemini Text: Multilingual Translation & Policy Generation", "We use AI to transcribe and translate local dialects into standard English for officials. Gemini is also used to generate strategic 'Mega Plans' by analyzing aggregated complaint data against municipal budgets.", G_GREEN)

    # ---------------------------------------------------------
    # SLIDE 5: WHO IT SERVES
    # ---------------------------------------------------------
    s5 = prs.slides.add_slide(blank_layout)
    set_bg(s5)
    add_header(s5, "Who It Serves", "Designed for the entire civic ecosystem", "05 / 12")
    
    add_feature_card(s5, Inches(1.0), Inches(2.5), Inches(3.5), Inches(3.0), "Citizens", "Empowers locals to report public infrastructure issues seamlessly in their native language and upvote community demands.", G_BLUE)
    add_feature_card(s5, Inches(4.9), Inches(2.5), Inches(3.5), Inches(3.0), "Municipal Crews", "Provides field teams with a prioritized, severity-ranked repair queue and direct Google Maps routing to incidents.", G_YELLOW)
    add_feature_card(s5, Inches(8.8), Inches(2.5), Inches(3.5), Inches(3.0), "National Policymakers", "Grants macro-level visibility into regional CapEx budget misalignment and generates high-level infrastructure investment directives.", G_RED)

    # ---------------------------------------------------------
    # SLIDE 6: CITIZEN FEATURES
    # ---------------------------------------------------------
    s6 = prs.slides.add_slide(blank_layout)
    set_bg(s6)
    add_header(s6, "App Feature: Citizen Gateway", "What we built for the public", "06 / 12")
    
    add_feature_card(s6, Inches(1.5), Inches(2.0), Inches(10.3), Inches(4.0), "Implemented Citizen Portal Features:", 
        "• GPS Auto-Detection: Instantly captures the precise location of the reported issue.\n"
        "• Multimodal Uploads: Citizens can upload photos of the defect and provide descriptions.\n"
        "• Community Demands Wall: A dedicated tab where citizens can propose and upvote macro-projects (e.g., new flyovers).\n"
        "• Live Tracking: A tracking hub to see if their submitted complaint is 'Under Review' or 'En Route'.\n"
        "• Themed UI: Fully responsive interface that works flawlessly on mobile devices.", G_BLUE)

    # ---------------------------------------------------------
    # SLIDE 7: OFFICIAL HUB FEATURES
    # ---------------------------------------------------------
    s7 = prs.slides.add_slide(blank_layout)
    set_bg(s7)
    add_header(s7, "App Feature: City Official Command Center", "What we built for government administrators", "07 / 12")
    
    add_feature_card(s7, Inches(1.5), Inches(2.0), Inches(10.3), Inches(4.0), "Implemented City Planner Features:", 
        "• Live Repair Queue: Automatically ingests citizen complaints and displays them with AI-assigned severity ratings.\n"
        "• CapEx Misalignment Matrix: A data table cross-referencing citizen demand volumes against planned municipal budgets.\n"
        "• Live Weather Telemetry: Integrated Open-Meteo API to display real-time weather conditions for the jurisdiction.\n"
        "• GIS Map Integration: Leaflet.js map integration for plotting incident coordinates.\n"
        "• Node Federation: Ability to toggle perspectives between different BRICS nations (India, Brazil, SA).", G_GREEN)

    # ---------------------------------------------------------
    # SLIDE 8: DEMO CITIZEN PORTAL
    # ---------------------------------------------------------
    s8 = prs.slides.add_slide(blank_layout)
    set_bg(s8)
    add_header(s8, "Demo: Citizen Gateway", "Screenshot of the Multimodal Citizen App", "08 / 12")
    
    add_screenshot_placeholder(s8, Inches(1.5), Inches(2.0), Inches(10.3), Inches(4.5), 
        "SCREENSHOT DEMO: CITIZEN PORTAL (INSERT HERE)", "http://localhost:8000/citizen")

    # ---------------------------------------------------------
    # SLIDE 9: DEMO CITY OFFICIAL HUB
    # ---------------------------------------------------------
    s9 = prs.slides.add_slide(blank_layout)
    set_bg(s9)
    add_header(s9, "Demo: City Official Command Center", "Screenshot of the AI Repair Queue & GIS Tools", "09 / 12")
    
    add_screenshot_placeholder(s9, Inches(1.5), Inches(2.0), Inches(10.3), Inches(4.5), 
        "SCREENSHOT DEMO: CITY OFFICIAL PORTAL (INSERT HERE)", "http://localhost:8000/city-official")

    # ---------------------------------------------------------
    # SLIDE 10: DEMO SOVEREIGN FEDERATION
    # ---------------------------------------------------------
    s10 = prs.slides.add_slide(blank_layout)
    set_bg(s10)
    add_header(s10, "Demo: Sovereign Node Federation", "Screenshot of Multi-Nation CapEx Overview", "10 / 12")
    
    add_screenshot_placeholder(s10, Inches(1.5), Inches(2.0), Inches(10.3), Inches(4.5), 
        "SCREENSHOT DEMO: SOVEREIGN NODE HUB (INSERT HERE)", "http://localhost:8000/central-official")

    # ---------------------------------------------------------
    # SLIDE 11: DEPLOYABILITY & SCALE
    # ---------------------------------------------------------
    s11 = prs.slides.add_slide(blank_layout)
    set_bg(s11)
    add_header(s11, "Deployability & How It Scales Across India", "Ready for real-world municipal adoption", "11 / 12")
    
    add_feature_card(s11, Inches(1.5), Inches(2.0), Inches(4.8), Inches(4.0), "Technical Foundation", 
        "• Zero-Setup Backend: Built on standard Python with a lightweight SQLite database, requiring no complex infrastructure to run.\n"
        "• Container Ready: Includes a Dockerfile and cloudbuild.yaml for 1-click deployment to Google Cloud Run.\n"
        "• API-First: Designed with standard webhook endpoints to easily integrate with WhatsApp/Telegram bots in the future.", G_BLUE)
        
    add_feature_card(s11, Inches(7.0), Inches(2.0), Inches(4.8), Inches(4.0), "Scaling Across India", 
        "• Localized Deployment: Any municipality can spin up their own 'node' of the system.\n"
        "• Low Latency AI: Uses Gemini Flash for instant triage of thousands of daily complaints.\n"
        "• Open Digital Public Good: Adheres to DPG principles, ensuring transparent data schemas that integrate with data.gov.in.", G_RED)

    # ---------------------------------------------------------
    # SLIDE 12: THANK YOU
    # ---------------------------------------------------------
    s12 = prs.slides.add_slide(blank_layout)
    set_bg(s12)
    
    tb = s12.shapes.add_textbox(Inches(1.0), Inches(1.5), Inches(11.333), Inches(1.5))
    tf = tb.text_frame
    p = tf.paragraphs[0]
    p.text = "Thank You!"
    p.font.name = 'Arial'
    p.font.size = Pt(48)
    p.font.bold = True
    p.font.color.rgb = G_BLUE
    p.alignment = PP_ALIGN.CENTER
    
    p2 = tf.add_paragraph()
    p2.text = "Team Syntax Error"
    p2.font.name = 'Arial'
    p2.font.size = Pt(24)
    p2.font.bold = True
    p2.font.color.rgb = TEXT_DARK
    p2.space_before = Pt(10)
    p2.alignment = PP_ALIGN.CENTER

    team = [
        ("Paras Singh", "Team Lead & Full-Stack", "Project concept, AI system design, and core platform development."),
        ("Aradhara Pal", "Frontend & UI/UX Designer", "Designed website layouts, mobile screens, and dashboards."),
        ("Anshika Anand", "Backend Developer", "Set up the server, database integration, and Render hosting."),
        ("Himanshi", "Product Research & Testing", "Researched city budget data, tested user workflows, and documentation.")
    ]
    
    for i, (name, role, desc) in enumerate(team):
        add_feature_card(s12, Inches(0.8 + i*3.0), Inches(3.8), Inches(2.8), Inches(2.8), name, f"{role}\n\n{desc}", [G_BLUE, G_RED, G_YELLOW, G_GREEN][i])

    out_name = "CivicPulse-BRICS_Pitch_Deck_Light.pptx"
    prs.save(out_name)
    print(f"Saved {out_name}")

if __name__ == "__main__":
    create_light_deck()

